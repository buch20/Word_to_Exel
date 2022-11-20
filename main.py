import docx
import os
import re
import openpyxl
from filter_list import *

path = 'C:\\Python_Project\\Parse_Word_to_Exel\\Медосмотр'
result = []
error = []


def returnpath(path_file):
    for name_file in os.listdir(path_file):
        if os.path.isdir(path_file + '\\' + name_file):
            returnpath(path_file + '\\' + name_file)
        elif os.path.isfile(path_file + '\\' + name_file):
            result.append(path_file + '\\' + name_file)

    return result


def reading_text(filename):
    doc = docx.Document(filename)
    text = []
    for parag in doc.paragraphs:
        text.append(parag.text)

    return '\n'.join(text)


def create_error_file(error_file):
    with open('error.txt', 'w', encoding='utf-8') as fe:
        fe.write(str('\n'.join(error_file)))


def main(path_file, count=0):
    book = openpyxl.Workbook()
    book.remove(book.active)
    for name in os.listdir(path_file):
        book.create_sheet(name)

    for file in returnpath(path_file):
        try:
            if file.split('.')[-1] == 'docx':
                read_file = reading_text(file)
                transform_file = (" ".join(read_file.split()))
                list_file = []
                for j in filter_list:
                    compile_filter = re.compile(j)
                    create_filter_file = compile_filter.search(str(transform_file)).group()
                    list_file.append(create_filter_file)
                for sheet in book.worksheets:
                    sheet['A1'] = "ФИО"
                    sheet['B1'] = "Класс"
                    sheet['C1'] = "Физическое развитие"
                    sheet['D1'] = "Группа здоровья"
                    sheet['E1'] = "Физ. группа"
                    sheet['F1'] = "Сопутствующие заболевания"
                    sheet.auto_filter.ref = 'A1:F999'
                    sheet.column_dimensions['A'].width = len(str(list_file[0]))
                    sheet.column_dimensions['B'].width = len("Класс  ")
                    sheet.column_dimensions['C'].width = len("Физическое развитие  ")
                    sheet.column_dimensions['D'].width = len("Группа здоровья  ")
                    sheet.column_dimensions['E'].width = len("Физ. группа  ")
                    sheet.column_dimensions['F'].width = len(str(list_file[5]))
                    if sheet.title == file.split('\\')[-2]:
                        sheet.append(list_file)

        except:
            if not file.split('\\')[-1] == 'шаблон медосмотр.docx':
                error.append(file)

        #
        finally:
            count += 1

    book.save(path.split('\\')[-1] + ".xlsx")
    book.close()

    create_error_file(error)


if __name__ == '__main__':
    main(path)
